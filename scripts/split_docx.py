#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""split_docx.py — 《AI战略传播力》Word 书稿 → VitePress Markdown 拆分管线（第一期导入工具）

输入：1-任总第一版/AI战略传播力2026-10-02.docx
输出：
  book/*.md                         34 个章节文件（00-封面 ~ 33-附录九）
  media/imageN.png                  55 张正文插图（按 body 流出现顺序编号）
  audits/2026-10-17-split-report.json  对账报告

设计要点（全部经源文档实测验证）：
  - 全文 2193 段均为 Normal 样式，无 Heading 层级 → 章节边界靠编号正则 + 全角空格（　）匹配
  - 章/篇/结语标题在「目录」与「正文」各出现一次 → 分组后取正文首次出现为锚点
    （陷阱：扉页导读里「第一章立题：…」也以「第一章」开头但无全角空格，正则已排除）
  - 篇扉页导读（【本篇导读】【主案例追踪】）归入各篇导论文件
  - 41 个 1×1 特殊框（通俗释义/本章案例/术语释义）转 Markdown 引用块，其余表格转 GFM 管道表
  - 图片按 body 流顺序导出，r:embed → rels.target_part 取真实媒体字节
  - 对账不达标（硬指标任一不符）以非零码退出

用法：python3 scripts/split_docx.py [--docx PATH] [--out DIR]
"""

import argparse
import json
import re
import sys
from pathlib import Path
from urllib.parse import quote

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

CN_NUM = "一二三四五六七八九十"
RE_CHAPTER = re.compile(rf"^第[{CN_NUM}]+章　")
RE_SECTION = re.compile(r"^(\d{1,2}\.\d{1,2})　(.+)$")
RE_PART = re.compile(rf"^第[{CN_NUM}]+篇　")
RE_PREFACE = re.compile(r"^前　　言$")
RE_CONCLUSION = re.compile(r"^结语　")
RE_APPENDIX = re.compile(rf"^附录[{CN_NUM}]+　")
RE_FIGURE_CAPTION = re.compile(r"^图\d{1,2}(?:-\d{1,2})?　")
RE_PART_MARK = re.compile(r"^【(本篇导读|主案例追踪[^】]*)】$")
RE_PART_RANGE = re.compile(r"^（第.+章）$")
RE_BYLINE = re.compile(r"^著　　者$|^2026年10月$")

PART_SHORT = {f"第{n}篇": f"第{n}篇导论" for n in "一二三四五"}
APPENDIX_SHORT = {
    "附录一": "四级跃迁传播力体检表", "附录二": "宣传系统AI传播能力训练清单",
    "附录三": "传播力评估指标体系", "附录四": "企业AI传播合规与内容审核清单",
    "附录五": "三主体协同传播模板", "附录六": "全球AI传播政策对照表",
    "附录七": "国家—城市—企业三级危机治理与信任修复工具包",
    "附录八": "出海与本地化传播工具包", "附录九": "测量指标词典与中文复测协议",
}


def fw2sp(text: str) -> str:
    """全角空格 → 半角空格（标题渲染用）。"""
    return text.replace("　", " ")


def render_runs(paragraph: Paragraph) -> str:
    """段落文本 + run 级加粗转换。"""
    parts = []
    for run in paragraph.runs:
        if not run.text:
            continue
        parts.append(f"**{run.text}**" if run.bold else run.text)
    out = "".join(parts).replace("****", "")
    return out.strip()


def table_to_gfm(table: Table) -> str:
    rows = []
    for row in table.rows:
        cells = [c.text.replace("\n", "<br>").replace("|", "\\|").strip() for c in row.cells]
        rows.append(cells)
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def box_to_quote(table: Table) -> str:
    """1×1 特殊框 → 引用块：首段标签加粗，其余行原样。"""
    lines = []
    for i, p in enumerate(table.rows[0].cells[0].paragraphs):
        text = p.text.strip()
        if not text:
            continue
        lines.append(f"> **{text}**" if i == 0 else f"> {text}")
    return "\n".join(lines)


def table_kind(table: Table) -> str:
    if len(table.rows) == 1 and len(table.columns) == 1:
        first = table.rows[0].cells[0].text.strip()
        if first.startswith("【通俗释义】"):
            return "box_paraphrase"
        if first.startswith("本章案例"):
            return "box_case"
        if first.startswith("术语释义") or first.startswith("【术语"):
            return "box_term"
        return "box_other"
    return "grid"


def chapter_short(label: str) -> str:
    """「第N章　短名——副题」→ 短名。"""
    name = label.split("　", 1)[1]
    return re.split(r"[——：]", name)[0].strip()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx", default="1-任总第一版/AI战略传播力2026-10-02.docx")
    ap.add_argument("--out", default="StratCom")
    args = ap.parse_args()

    docx_path = Path(args.docx).resolve()
    out_root = Path(args.out).resolve()
    book_dir, media_dir, audits_dir = out_root / "book", out_root / "media", out_root / "audits"
    for d in (book_dir, media_dir, audits_dir):
        d.mkdir(parents=True, exist_ok=True)

    doc = Document(str(docx_path))
    body = doc.element.body

    # ---------- 阶段 1：body 流序列（段落与表格交错，不丢表） ----------
    blocks, para_texts = [], []
    for ch in body:
        tag = ch.tag.split("}")[1]
        if tag == "p":
            p = Paragraph(ch, doc)
            blocks.append({"kind": "p", "obj": p})
            para_texts.append(p.text.strip())
        elif tag == "tbl":
            blocks.append({"kind": "tbl", "obj": Table(ch, doc)})
    n_paras = len(para_texts)

    # ---------- 阶段 2：锚点定位（分组取正文首次出现） ----------
    toc_start = next(i for i, t in enumerate(para_texts) if t == "目　录")
    assert toc_start == 15, f"目录起点异常: {toc_start}"

    preface_hits = [i for i, t in enumerate(para_texts) if RE_PREFACE.match(t)]
    assert len(preface_hits) == 1, f"「前　　言」应恰好出现 1 次，实际 {preface_hits}"
    preface_anchor = preface_hits[0]
    assert para_texts[preface_anchor + 1], "前言副题段为空，H1 拼接会缺副题"
    toc_end = preface_anchor - 1  # 目录区 = [toc_start, toc_end]

    def anchors_by_key(pattern, name):
        groups = {}
        for i, t in enumerate(para_texts):
            if pattern.match(t):
                groups.setdefault(t.split("　")[0], []).append(i)
        out = {}
        for key, idxs in groups.items():
            after = [i for i in idxs if i > toc_end]
            assert after, f"{name}「{key}」未找到正文锚点"
            out[key] = after[0]
        return out

    chapters = anchors_by_key(RE_CHAPTER, "章")
    parts = anchors_by_key(RE_PART, "篇")
    conclusions = anchors_by_key(RE_CONCLUSION, "结语")
    appendix_hits = [i for i, t in enumerate(para_texts) if RE_APPENDIX.match(t)]

    assert len(chapters) == 17, f"章锚点 {len(chapters)} != 17"
    assert len(parts) == 5, f"篇锚点 {len(parts)} != 5"
    assert len(conclusions) == 1, f"结语锚点 {len(conclusions)} != 1"
    assert len(appendix_hits) == 9, f"附录锚点 {len(appendix_hits)} != 9"

    anchors = [(preface_anchor, "preface", para_texts[preface_anchor])]
    anchors += [(i, "part", para_texts[i]) for i in parts.values()]
    anchors += [(i, "chapter", para_texts[i]) for i in chapters.values()]
    anchors += [(i, "conclusion", para_texts[i]) for i in conclusions.values()]
    anchors += [(i, "appendix", para_texts[i]) for i in appendix_hits]
    anchors.sort(key=lambda x: x[0])
    assert len(anchors) == 33, f"锚点总数 {len(anchors)} != 33"

    spans = []
    for idx, (a, kind, label) in enumerate(anchors):
        end = anchors[idx + 1][0] - 1 if idx + 1 < len(anchors) else n_paras - 1
        spans.append({"start": a, "end": end, "kind": kind, "label": label})

    # ---------- 阶段 3：图片导出（按 body 流顺序） ----------
    media_by_pidx = {}   # pidx -> [(media_name, orig_target)]
    img_counter = 0
    pidx = -1
    for b in blocks:
        if b["kind"] != "p":
            continue
        pidx += 1
        blips = b["obj"]._p.findall(".//" + qn("a:blip"))
        if not blips:
            continue
        entries = []
        for blip in blips:
            rid = blip.get(qn("r:embed"))
            rel = doc.part.rels[rid]
            img_counter += 1
            name = f"image{img_counter}.png"
            (media_dir / name).write_bytes(rel.target_part.blob)
            entries.append((name, rel.target_ref))
        media_by_pidx[pidx] = entries

    fig_caption_paras = sum(1 for t in para_texts if RE_FIGURE_CAPTION.match(t))
    assert img_counter == fig_caption_paras == 55, \
        f"图片对账失败: blip={img_counter} 图题={fig_caption_paras}（应均为 55）"

    # ---------- 阶段 4：逐 section 渲染 ----------
    def render_span(span, filename):
        lines = []
        label = fw2sp(span["label"])
        if span["kind"] == "preface":
            subtitle = fw2sp(para_texts[span["start"] + 1])
            lines.append(f"# 前言 {subtitle}")
        else:
            lines.append(f"# {label}")

        n_img = n_grid = n_box = 0
        char_count = 0
        refs_found = 0
        pidx = -1
        for b in blocks:
            if b["kind"] == "p":
                pidx += 1
                if not (span["start"] <= pidx <= span["end"]):
                    continue
                t = para_texts[pidx]
                if span["kind"] == "preface" and pidx in (span["start"], span["start"] + 1):
                    char_count += len(t)
                    continue  # 扉页段与副题段已并入 H1
                if pidx == span["start"]:
                    char_count += len(t)
                    continue  # 锚点段已渲染为 H1，跳过原文重复
                if pidx in media_by_pidx:
                    for name, _ in media_by_pidx[pidx]:
                        lines.append(f"![](../media/{name})")
                        n_img += 1
                    continue
                if not t:
                    continue
                if span["kind"] == "part" and RE_PART_MARK.match(t):
                    lines.append(f"**{t}**")
                elif span["kind"] == "part" and RE_PART_RANGE.match(t):
                    lines.append(t)
                elif t == "本章参考文献":
                    lines.append("## 本章参考文献")
                    refs_found += 1
                elif RE_SECTION.match(t):
                    m = RE_SECTION.match(t)
                    lines.append(f"## {m.group(1)} {m.group(2)}")
                elif span["kind"] in ("preface", "conclusion", "appendix") and \
                        re.match(rf"^[{CN_NUM}]+、", t):
                    lines.append(f"## {t}")
                else:
                    lines.append(render_runs(b["obj"]))  # 图题/文献行/著者/正文统一原样
                char_count += len(t)
            else:
                if not (span["start"] <= pidx <= span["end"]):
                    continue
                tb = b["obj"]
                kind = table_kind(tb)
                if kind.startswith("box"):
                    lines += ["", box_to_quote(tb), ""]
                    n_box += 1
                else:
                    lines += ["", table_to_gfm(tb), ""]
                    n_grid += 1
                for row in tb.rows:
                    for cell in row.cells:
                        char_count += len(cell.text)

        text = re.sub(r"\n{3,}", "\n\n", "\n\n".join(lines)).strip() + "\n"
        (book_dir / filename).write_text(text, encoding="utf-8")
        return {"file": filename, "chars": char_count, "images": n_img,
                "grid_tables": n_grid, "box_tables": n_box, "refs": refs_found}

    # 文件命名表（00-33）
    file_plan = [(None, "00-封面.md", "封面")]
    for no, span in enumerate(spans, start=1):
        kind = span["kind"]
        if kind == "preface":
            fn = "01-前言 AI时代必须重建战略传播.md"
        elif kind == "part":
            fn = f"{no:02d}-{PART_SHORT[span['label'].split('　')[0]]}.md"
        elif kind == "chapter":
            fn = f"{no:02d}-{chapter_short(span['label'])}.md"
        elif kind == "conclusion":
            fn = f"{no:02d}-结语 传播力即竞争力.md"
        else:
            key = span["label"].split("　")[0]
            fn = f"{no:02d}-{key} {APPENDIX_SHORT[key]}.md"
        file_plan.append((span, fn, span["label"]))
    assert len(file_plan) == 34, f"文件数 {len(file_plan)} != 34"
    assert len({fn for _, fn, _ in file_plan}) == 34, "文件名重复"

    # 00-封面.md（含全书目录链接）
    cover = ["# AI战略传播力", "",
             "**AI时代的战略传播理论体系与行动框架**", "",
             "**钟诚　任国刚　吴鹏　著**", "", "新华出版社", "", "## 全书目录", ""]
    for span, fn, label in file_plan[1:]:
        disp = "前言 AI时代必须重建战略传播" if span["kind"] == "preface" else fw2sp(label)
        cover.append(f"- [{disp}]({quote(fn[:-3])})")
    (book_dir / "00-封面.md").write_text("\n".join(cover) + "\n", encoding="utf-8")

    # 渲染 33 个正文 section
    results = [render_span(span, fn) for span, fn, _ in file_plan if span]

    # ---------- 阶段 5：对账 ----------
    # 源字符基线：段落文本 + 表格 cell 文本；扣除封面与目录区（段 0 .. toc_end）
    src_chars = sum(len(t) for i, t in enumerate(para_texts) if i > toc_end and t)
    for tb in doc.tables:
        for row in tb.rows:
            for cell in row.cells:
                src_chars += len(cell.text)

    out_chars = sum(r["chars"] for r in results)  # 不含封面（封面为新建内容）
    box_kinds = {"box_paraphrase": 0, "box_case": 0, "box_term": 0, "box_other": 0}
    grid_total = 0
    for tb in doc.tables:
        k = table_kind(tb)
        if k == "grid":
            grid_total += 1
        else:
            box_kinds[k] += 1
    refs_total = sum(r["refs"] for r in results)
    deviation = abs(out_chars - src_chars) / src_chars * 100

    checks = [
        ("章节文件数", len(file_plan), 34),
        ("插图数(blip)", img_counter, 55),
        ("图题段数", fig_caption_paras, 55),
        ("网格表数", grid_total, 131),
        ("1x1特殊框数", sum(box_kinds.values()), 41),
        ("表格总数", grid_total + sum(box_kinds.values()), 172),
        ("本章参考文献", refs_total, 17),
    ]
    ok = (len(file_plan) == 34 and img_counter == 55 and fig_caption_paras == 55
          and grid_total == 131 and sum(box_kinds.values()) == 41
          and refs_total == 17 and deviation < 1.0)

    report = {
        "source": str(docx_path),
        "generated_by": "scripts/split_docx.py",
        "files": results,
        "media": {
            "images": img_counter,
            "naming": "按 body 流出现顺序 image1-55.png",
            "map": {str(k): v for k, v in media_by_pidx.items()},
        },
        "tables": {"grid": grid_total, "boxes": box_kinds},
        "references_sections": refs_total,
        "char_accounting": {"source_chars": src_chars, "output_chars": out_chars,
                            "deviation_pct": round(deviation, 3)},
        "checks": [{"item": n, "got": g, "expect": w} for n, g, w in checks],
        "passed": ok,
    }
    (audits_dir / "2026-10-17-split-report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"{'文件':<46}{'字符':>8}{'图':>4}{'网格表':>6}{'框':>4}{'文献':>4}")
    for r in results:
        print(f"{r['file']:<46}{r['chars']:>8}{r['images']:>4}{r['grid_tables']:>6}{r['box_tables']:>4}{r['refs']:>4}")
    print()
    for name, got, want in checks:
        print(f"  {name}: {got} (期望 {want})")
    print(f"\n源字符 {src_chars} / 输出字符 {out_chars} / 偏差 {round(deviation, 3)}%")
    print(f"1x1 框分类: {box_kinds}")
    print("对账结果:", "PASS" if ok else "FAIL")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
